"""Convert the organisers' constraints workbook into ``config/operational_limits.yaml``.

Run once (and again whenever the workbook changes)::

    uv run --with openpyxl python scripts/import_constraints_workbook.py

``openpyxl`` is deliberately *not* a runtime dependency: the committed YAML is
the source the simulator reads, and the workbook SHA-256 in its header proves
which workbook it came from.  Output is deterministic (workbook row order, no
timestamps) so re-running on the same workbook produces an identical file.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path
from typing import Any

import yaml

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "challenge-2-operational-constraints.xlsx"
DEFAULT_OUT = ROOT / "config" / "operational_limits.yaml"

SHEETS = ["airport_capacity", "wake_separation", "fuel_and_crew", "rules_of_play"]
AIRPORT_HEADER_ROW = 3  # 1-based; rows 1-2 hold the sheet note

AIRPORT_COLUMNS = [
    "IATA_CODE",
    "AIRPORT",
    "CITY",
    "STATE",
    "delayed_departures_in_flights_csv",
    "size_class",
    "active_runways",
    "gates",
    "max_runway_ops_per_hour",
    "min_turnaround_min",
    "spare_gates_at_any_time",
]
INT_COLUMNS = {
    "delayed_departures_in_flights_csv",
    "active_runways",
    "gates",
    "max_runway_ops_per_hour",
    "min_turnaround_min",
    "spare_gates_at_any_time",
}

# Workbook wake classes -> ICAO wake categories used in the code.
WAKE_CLASS = {"Super": "J", "Heavy": "H", "Large": "M", "Small": "L"}

# fuel_and_crew rule name -> (key, parser) ; parser(text) -> dict of numbers.
_NUM = r"(\d[\d,]*(?:\.\d+)?)"


def _numbers(text: str) -> list[float]:
    return [float(m.replace(",", "")) for m in re.findall(_NUM, text)]


def _first(text: str) -> float:
    nums = _numbers(text)
    if not nums:
        raise ValueError(f"no number in {text!r}")
    return nums[0]


def _as_int(x: float) -> int | float:
    return int(x) if float(x).is_integer() else x


FUEL_CREW_RULES: dict[str, tuple[str, Any]] = {
    "Final fuel reserve": (
        "final_reserve",
        lambda t: {"holding_min": _as_int(_first(t))},
    ),
    "Contingency fuel": (
        "contingency",
        lambda t: {"fraction_of_trip_fuel": round(_first(t) / 100.0, 6)},
    ),
    "Holding fuel burn (narrowbody)": (
        "holding_burn_narrowbody",
        lambda t: {"kg_per_hour": _as_int(_first(t))},
    ),
    "Holding fuel burn (widebody)": (
        "holding_burn_widebody",
        lambda t: {"kg_per_hour": _as_int(_first(t))},
    ),
    "Assumed spare fuel on arrival": (
        "spare_fuel_on_arrival",
        lambda t: {"holding_min": _as_int(_first(t))},
    ),
    "Flight crew max duty": (
        "flight_crew_max_duty",
        lambda t: {"hours": _as_int(_first(t)), "minutes": _as_int(_first(t) * 60)},
    ),
    "Ground crew shift": (
        "ground_crew_shift",
        lambda t: {
            "hours": _as_int(_numbers(t)[0]),
            "minutes": _as_int(_numbers(t)[0] * 60),
            "overtime_multiplier": _as_int(_numbers(t)[1]),
        },
    ),
    "Gate change notice": (
        "gate_change_notice",
        lambda t: {"min_before_arrival_min": _as_int(_first(t))},
    ),
}


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _sheet_rows(ws: Any) -> list[tuple[Any, ...]]:
    return [tuple(_clean(v) for v in row) for row in ws.iter_rows(values_only=True)]


def _airport_capacity(ws: Any) -> dict[str, Any]:
    rows = _sheet_rows(ws)
    note = rows[0][0]
    header = [str(h) for h in rows[AIRPORT_HEADER_ROW - 1] if h is not None]
    if header != AIRPORT_COLUMNS:
        raise ValueError(f"unexpected airport_capacity header: {header}")
    airports: dict[str, dict[str, Any]] = {}
    for row in rows[AIRPORT_HEADER_ROW:]:
        if row[0] is None:
            continue
        record = dict(zip(AIRPORT_COLUMNS, row))
        code = str(record.pop("IATA_CODE"))
        for col in INT_COLUMNS:
            record[col] = int(record[col])
        if code in airports:
            raise ValueError(f"duplicate airport {code}")
        airports[code] = record  # original workbook column names, minus the key
    return {
        "note": note,
        "header_row": AIRPORT_HEADER_ROW,
        "columns": AIRPORT_COLUMNS,
        "airports": airports,
    }


def _wake_class(label: str) -> str:
    word = label.split()[0]
    if word not in WAKE_CLASS:
        raise ValueError(f"unknown wake class {label!r}")
    return WAKE_CLASS[word]


def _wake_separation(ws: Any) -> dict[str, Any]:
    rows = _sheet_rows(ws)
    note = rows[0][0]
    pairs: list[dict[str, Any]] = []
    other: dict[str, Any] | None = None
    for row in rows[3:]:
        leader, follower, nm, secs = (list(row) + [None] * 4)[:4]
        if leader is None:
            continue
        if str(leader).lower().startswith("any other"):
            other = {"label": leader, "min_separation_nm": nm, "min_time_s": secs}
            continue
        pairs.append(
            {
                "leader": leader,
                "follower": follower,
                "leader_cat": _wake_class(str(leader)),
                "follower_cat": _wake_class(str(follower)),
                "min_separation_nm": nm,
                "min_time_s": secs,
            }
        )
    if other is None:
        raise ValueError("wake_separation: missing 'Any other pair' row")
    return {
        "note": note,
        "class_map": {f"{k}": v for k, v in WAKE_CLASS.items()},
        "pairs": pairs,
        "any_other_pair": other,
    }


def _fuel_and_crew(ws: Any) -> dict[str, Any]:
    rows = _sheet_rows(ws)
    out: dict[str, Any] = {}
    for row in rows[1:]:
        rule, value, applies_to = (list(row) + [None] * 3)[:3]
        if rule is None:
            continue
        if rule not in FUEL_CREW_RULES:
            raise ValueError(f"unknown fuel_and_crew rule {rule!r}")
        key, parse = FUEL_CREW_RULES[rule]
        out[key] = {
            "rule": rule,
            "text": value,
            "applies_to": applies_to,
            **parse(str(value)),
        }
    missing = {k for k, _ in FUEL_CREW_RULES.values()} - set(out)
    if missing:
        raise ValueError(f"fuel_and_crew missing rules: {sorted(missing)}")
    return out


def _rules_of_play(ws: Any) -> dict[str, str]:
    rows = _sheet_rows(ws)
    return {str(r[0]): str(r[1]) for r in rows[1:] if r and r[0] is not None}


def build(xlsx: Path) -> tuple[str, dict[str, Any]]:
    import openpyxl  # imported lazily: not a runtime dependency

    digest = hashlib.sha256(xlsx.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    if wb.sheetnames != SHEETS:
        raise ValueError(f"unexpected sheets {wb.sheetnames}; expected {SHEETS}")
    data: dict[str, Any] = {
        "source": {
            "workbook": xlsx.name,
            "workbook_sha256": digest,
            "sheets": list(wb.sheetnames),
            "generated_by": "scripts/import_constraints_workbook.py",
        },
        "airport_capacity": _airport_capacity(wb["airport_capacity"]),
        "wake_separation": _wake_separation(wb["wake_separation"]),
        "fuel_and_crew": _fuel_and_crew(wb["fuel_and_crew"]),
        "rules_of_play": _rules_of_play(wb["rules_of_play"]),
    }
    wb.close()
    return digest, data


def render(digest: str, data: dict[str, Any], xlsx_name: str) -> str:
    header = (
        "# GENERATED by scripts/import_constraints_workbook.py -- do not edit by hand.\n"
        f"# Source workbook: {xlsx_name}\n"
        f"# Workbook sha256: {digest}\n"
        f"# Sheets: {', '.join(data['source']['sheets'])}\n"
        "# Deliberate deviations from these limits live in config/limits_deviations.yaml.\n"
    )
    body = yaml.safe_dump(
        data, sort_keys=False, allow_unicode=True, width=1000, default_flow_style=None
    )
    return header + body


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args(argv)

    digest, data = build(args.xlsx)
    text = render(digest, data, args.xlsx.name)
    args.out.write_text(text, encoding="utf-8")
    n = len(data["airport_capacity"]["airports"])
    print(f"wrote {args.out.relative_to(ROOT) if args.out.is_relative_to(ROOT) else args.out}: "
          f"{n} airports, {len(data['wake_separation']['pairs'])} wake pairs, sha256 {digest[:12]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
